import io

import pymupdf
from docx import Document


def make_pdf(pages: list[list[tuple[str, float]]], *, password: str | None = None) -> bytes:
    """One page per entry; each line is (text, font size)."""
    pdf = pymupdf.open()
    for lines in pages:
        page = pdf.new_page()
        y = 72.0
        for text, size in lines:
            page.insert_text((72, y), text, fontsize=size)
            y += size * 1.8
    if password:
        encrypted = pdf.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password + "x"
        )
        return bytes(encrypted)
    return bytes(pdf.tobytes())


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Leave policy", level=1)
    doc.add_paragraph("Employees get 25 days of paid leave.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tier"
    table.cell(0, 1).text = "Days"
    table.cell(1, 0).text = "Senior"
    table.cell(1, 1).text = "30"
    doc.add_paragraph("Ask HR for details.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
