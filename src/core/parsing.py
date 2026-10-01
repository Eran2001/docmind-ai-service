import io
from collections import Counter
from collections.abc import Callable
from typing import Any

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from core.cleaning import ParsedDocument, ParsedPage
from core.errors import AppError, ErrorCode
from core.logging import get_logger

log = get_logger(__name__)

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT = "text/plain"
MARKDOWN = "text/markdown"

_HEADING_SIZE_RATIO = 1.15
_HEADING_MAX_CHARS = 100
_BOLD_FLAG = 16


def parse_file(data: bytes, mime_type: str) -> ParsedDocument:
    """Extracts text from a PDF, DOCX, TXT or MD file. Headings come out as Markdown (`# Title`) lines."""
    parser = _PARSERS.get(mime_type)
    if parser is None:
        raise AppError(ErrorCode.PARSE_FAILED, f"Unsupported file type: {mime_type}.")
    document = parser(data)

    if not any(page.text.strip() for page in document.pages):
        raise AppError(ErrorCode.EMPTY_DOCUMENT, "No text could be found in this document.")
    return document


def _parse_text(data: bytes) -> ParsedDocument:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise AppError(ErrorCode.PARSE_FAILED, "The file is not valid UTF-8 text.") from None
    return ParsedDocument(pages=[ParsedPage(page_number=None, text=text.replace("\x00", ""))])


def _parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = Document(io.BytesIO(data))
        blocks: list[str] = []
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if not text:
                    continue
                style = item.style.name if item.style and item.style.name else ""
                is_heading = style == "Title" or style.startswith("Heading")
                blocks.append(f"# {text}" if is_heading else text)
            elif isinstance(item, Table):
                blocks.extend(_table_rows(item))
    except Exception as exc:
        log.warning("docx_parse_failed", error_type=type(exc).__name__)
        raise AppError(ErrorCode.PARSE_FAILED, "This Word file could not be read.") from exc
    return ParsedDocument(pages=[ParsedPage(page_number=None, text="\n\n".join(blocks))])


def _table_rows(table: Table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            text = " ".join(cell.text.split())
            if text and (not cells or cells[-1] != text):  # merged cells repeat the same text
                cells.append(text)
        if cells:
            rows.append(" | ".join(cells))
    return rows


def _parse_pdf(data: bytes) -> ParsedDocument:
    try:
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise AppError(ErrorCode.PARSE_FAILED, "This PDF is password-protected.")
            page_blocks: list[list[dict[str, Any]]] = [page.get_text("dict")["blocks"] for page in pdf]
    except AppError:
        raise
    except Exception as exc:
        log.warning("pdf_parse_failed", error_type=type(exc).__name__)
        raise AppError(ErrorCode.PARSE_FAILED, "This PDF could not be read.") from exc

    body_size = _body_font_size(page_blocks)
    pages = [
        ParsedPage(page_number=number, text=_pdf_page_text(blocks, body_size))
        for number, blocks in enumerate(page_blocks, start=1)
    ]
    return ParsedDocument(pages=pages, page_count=len(pages))


def _text_blocks(blocks: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    """Per text block, its non-empty lines as {text, size, bold}."""
    result: list[list[dict[str, Any]]] = []
    for block in blocks:
        if block.get("type") != 0:
            continue
        lines: list[dict[str, Any]] = []
        for line in block["lines"]:
            spans = line["spans"]
            text = "".join(span["text"] for span in spans).replace("\x00", "").strip()
            if text:
                lines.append(
                    {
                        "text": text,
                        "size": max(span["size"] for span in spans),
                        "bold": all(span["flags"] & _BOLD_FLAG for span in spans),
                    }
                )
        if lines:
            result.append(lines)
    return result


def _body_font_size(page_blocks: list[list[dict[str, Any]]]) -> float:
    """The font size most characters are set in: the body text."""
    weights: Counter[float] = Counter()
    for blocks in page_blocks:
        for lines in _text_blocks(blocks):
            for line in lines:
                weights[round(line["size"], 1)] += len(line["text"])
    return weights.most_common(1)[0][0] if weights else 0.0


def _pdf_page_text(blocks: list[dict[str, Any]], body_size: float) -> str:
    paragraphs: list[str] = []
    for lines in _text_blocks(blocks):
        out: list[str] = []
        for line in lines:
            text: str = line["text"]
            short = len(text) <= _HEADING_MAX_CHARS and not text.endswith(".")
            larger = body_size > 0 and line["size"] >= body_size * _HEADING_SIZE_RATIO
            lone_bold = line["bold"] and len(lines) == 1 and len(text) <= _HEADING_MAX_CHARS // 2
            out.append(f"# {text}" if short and (larger or lone_bold) else text)
        paragraphs.append("\n".join(out))
    return "\n\n".join(paragraphs)


_PARSERS: dict[str, Callable[[bytes], ParsedDocument]] = {
    PDF: _parse_pdf,
    DOCX: _parse_docx,
    TEXT: _parse_text,
    MARKDOWN: _parse_text,
}
